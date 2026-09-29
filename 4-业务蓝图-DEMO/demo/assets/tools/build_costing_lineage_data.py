#!/usr/bin/env python3
"""从批次样例工作簿构建离线 Demo 数据。

该脚本只生成批次与数量关系，不读取、计算或输出成本字段。
"""

from __future__ import annotations

import json
import re
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
DEMO_DIR = HERE.parents[1]
SOURCE_DIR = DEMO_DIR.parents[2] / "成本分摊模块" / "批次划分"
REL_BOOK = SOURCE_DIR / "批次和批次关联数据样例_QC整罐转成品罐修正版.xlsx"
OUTPUT_JSON = DEMO_DIR / "assets" / "data" / "costing" / "costing-lineage-data.json"
OUTPUT_JS = DEMO_DIR / "assets" / "scripts" / "costing" / "costing-lineage-data.js"


def fmt(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if value is None:
        return "—"
    if isinstance(value, float):
        return round(value, 3)
    return value


def pct(value):
    if value in (None, "", "—"):
        return "—"
    value = float(value)
    if value <= 1.000001:
        value *= 100
    return f"{value:.1f}%"


batches = {}
relations = []


def add_batch(
    batch_id,
    *,
    order,
    stage,
    material,
    device,
    start,
    end,
    qty,
    unit="吨",
    line="果葡糖浆",
    route="果葡糖浆产线",
    status="已生成",
    rule="系统按已生效批次划分规则生成",
    source="Excel样例",
    note="",
):
    if not batch_id:
        return
    item = {
        "id": str(batch_id),
        "period": "2026-08",
        "line": line,
        "route": route,
        "order": int(order),
        "stage": stage,
        "material": material,
        "device": device or "—",
        "start": fmt(start),
        "end": fmt(end),
        "qty": round(float(qty or 0), 3),
        "unit": unit,
        "status": status,
        "rule": rule,
        "source": source,
        "note": note,
    }
    if batch_id in batches:
        current = batches[batch_id]
        for key, value in item.items():
            if current.get(key) in (None, "", "—", 0) and value not in (None, "", "—", 0):
                current[key] = value
        return
    batches[batch_id] = item


def add_relation(
    rel_id,
    upstream,
    downstream,
    *,
    layer,
    overlap="—",
    qty=0,
    upstream_share="—",
    downstream_share="—",
    yield_rate="—",
    method="时间窗重叠匹配",
    evidence="DCS时间窗与生产实绩",
    confidence="B",
    status="系统确认",
    source="Excel样例",
):
    relations.append(
        {
            "id": str(rel_id),
            "from": str(upstream),
            "to": str(downstream),
            "layer": layer,
            "overlap": fmt(overlap),
            "qty": round(float(qty or 0), 3),
            "unit": "吨",
            "upstreamShare": pct(upstream_share),
            "downstreamShare": pct(downstream_share),
            "yield": pct(yield_rate),
            "method": method or "时间窗重叠匹配",
            "evidence": evidence,
            "confidence": confidence,
            "status": status,
            "source": source,
        }
    )


def rows(ws, start, end):
    return [[cell.value for cell in row] for row in ws.iter_rows(min_row=start, max_row=end)]


wb = load_workbook(REL_BOOK, read_only=True, data_only=True)

# 2→3：线边库、关系、投料段
ws = wb["2-3线边库到投料批次"]
for r in rows(ws, 4, 63):
    add_batch(
        r[0], order=2, stage="线边库批次", material="玉米淀粉", device="XB01线边工位",
        start=r[1], end=r[2], qty=r[4],
        rule="指定移库叉车与线边工位，按工位及移库作业划分",
        source="Excel·2-3线边库到投料批次",
    )
for r in rows(ws, 130, 144):
    add_batch(
        r[0], order=3, stage="投料批次", material="淀粉乳", device="投料口—淀粉乳接收罐",
        start=r[1], end=r[2], qty=r[8],
        rule="按投料口至糖化罐入口的动态过程时间窗 Tw 切片",
        source="Excel·2-3线边库到投料批次",
    )
for r in rows(ws, 67, 126):
    add_relation(
        f"REL-F55-23-{int(r[0]):03d}", r[2], r[1], layer="线边库→投料",
        overlap=r[3], qty=r[6], upstream_share=r[7], downstream_share=r[8], yield_rate=r[5],
        method=r[9], evidence="线边投料记录+投料段产出折算", confidence="A",
        source="Excel·2-3线边库到投料批次",
    )

# 3→4：投料段→糖化罐
ws = wb["3-4投料和糖化批次"]
for r in rows(ws, 49, 60):
    add_batch(
        r[0], order=4, stage="糖化批次", material="糖化液", device=r[1],
        start=r[2], end=r[7], qty=r[9], status="已出罐",
        rule="12个糖化罐按独立进出罐周期划分，以糖化罐出口阀为物理切分点",
        source="Excel·3-4投料和糖化批次",
    )
for r in rows(ws, 22, 45):
    add_relation(
        f"REL-F55-34-{int(r[0]):03d}", r[2], r[1], layer="投料→糖化",
        overlap=r[3], qty=r[4], upstream_share=r[5], downstream_share=r[6],
        method=r[7], evidence="糖化罐进料开始/结束时间与投料段有效窗", confidence="B",
        source="Excel·3-4投料和糖化批次",
    )

# 4→5：糖化罐→F42中间品
ws = wb["4-5糖化和F42批次"]
for r in rows(ws, 36, 43):
    add_batch(
        r[0], order=5, stage="F42中间品批次", material="F42中间品", device="F42装置·FT120101",
        start=r[1], end=r[2], qty=r[8],
        rule="按糖化罐出口至离交出料流量计 FT120101 的动态时间窗 Tf42 切片",
        source="Excel·4-5糖化和F42批次",
    )
for r in rows(ws, 19, 32):
    add_relation(
        f"REL-F55-45-{int(r[0]):03d}", r[2], r[1], layer="糖化→F42",
        overlap=r[3], qty=r[6], upstream_share=r[7], downstream_share=r[8], yield_rate=r[5],
        method=r[9], evidence="糖化罐出料窗×恒定流量", confidence="B",
        source="Excel·4-5糖化和F42批次",
    )

# 5→6：F42→色谱分离（F90）
ws = wb["5-6F42和色谱批次"]
for r in rows(ws, 34, 44):
    add_batch(
        r[0], order=6, stage="色谱分离批次", material="F90糖浆", device="色谱分离·FT120103",
        start=r[1], end=r[2], qty=r[8],
        rule="按离交出料流量计 FT120103 至 F55 混合器调节阀的动态时间窗 Tf90 切片",
        source="Excel·5-6F42和色谱批次",
    )
for r in rows(ws, 15, 30):
    add_relation(
        f"REL-F55-56-{int(r[0]):03d}", r[2], r[1], layer="F42→色谱",
        overlap=r[3], qty=r[6], upstream_share=r[7], downstream_share=r[8], yield_rate=r[5],
        method=r[9], evidence="F42有效窗+色谱周期+累计流量", confidence="B",
        source="Excel·5-6F42和色谱批次",
    )

# 5/6→7：F42直供与F90汇合为F55
ws = wb["5和6-7汇合到F55批次"]
for r in rows(ws, 64, 74):
    add_batch(
        r[0], order=7, stage="F55混合批次", material="果糖F55", device="F55混合器—蒸发器",
        start=r[1], end=r[10], qty=r[8],
        rule="F42直供与F90按配方同步汇合；配方切换或动态时间窗 Th 结束时生成新批次",
        source="Excel·5和6-7汇合到F55批次",
        note=f"F42:F90={r[5]}",
    )
for r in rows(ws, 34, 60):
    add_relation(
        f"REL-F55-67-{int(r[0]):03d}", r[2], r[1], layer="F42/F90→F55",
        overlap=r[3], qty=r[6], upstream_share=r[7], downstream_share=r[8], yield_rate=r[5],
        method=r[9], evidence="混合器阀门同步+配方执行记录", confidence="A",
        source="Excel·5和6-7汇合到F55批次",
    )

# 7→8：F55→QC罐
ws = wb["7-8F55和QC罐批次"]
for r in rows(ws, 77, 122):
    add_batch(
        r[0], order=8, stage="QC罐批次", material="果糖F55", device=r[1],
        start=r[2], end=r[9], qty=r[10], status="已放行" if r[6] == "是" else "末批未满·已放行",
        rule="QC罐独立进料，以进料阀和出料阀为物理边界，检验放行后可转罐",
        source="Excel·7-8F55和QC罐批次",
    )
for r in rows(ws, 18, 73):
    add_relation(
        f"REL-F55-78-{int(r[0]):03d}", r[2], r[1], layer="F55→QC罐",
        overlap=f"{fmt(r[3])} 至 {fmt(r[4])}", qty=r[7], upstream_share=r[8], downstream_share=r[9], yield_rate=r[6],
        method="QC罐进料窗×累计流量", evidence="DCS进罐阀状态+进罐累计量", confidence="A",
        source="Excel·7-8F55和QC罐批次",
    )

# 8→9：QC整罐转成品罐（不得拆分）
ws = wb["8-9QC罐和成品罐批次"]
for r in rows(ws, 102, 107):
    add_batch(
        r[0], order=9, stage="成品罐批次", material="果糖F55", device=r[1],
        start=r[2], end=r[3], qty=r[9], status="已入库",
        rule="按产品分罐并按进罐时间切分；QC批次必须整罐转入同一成品罐",
        source="Excel·8-9QC罐和成品罐批次",
        note=r[7],
    )
for r in rows(ws, 53, 98):
    add_relation(
        f"REL-F55-89-{int(r[0]):03d}", r[2], r[1], layer="QC罐→成品罐",
        overlap=f"{fmt(r[3])} 至 {fmt(r[4])}", qty=r[7], upstream_share=r[8], downstream_share=r[9], yield_rate=r[6],
        method=r[10], evidence="QC放行记录+整罐转罐作业记录", confidence="A",
        source="Excel·8-9QC罐和成品罐批次",
    )

# 9→10：成品罐→装车
ws = wb["9-10成品罐和装车批次"]
for r in rows(ws, 77, 136):
    add_batch(
        r[0], order=10, stage="车次批次", material="果糖F55", device="成品装车位",
        start=r[1], end=r[2], qty=r[3], status="已装车",
        rule="每车独立成批，以成品罐出料阀为起点并绑定车号",
        source="Excel·9-10成品罐和装车批次",
    )
for r in rows(ws, 13, 73):
    add_relation(
        f"REL-F55-910-{int(r[0]):03d}", r[2], r[1], layer="成品罐→装车",
        qty=r[3], upstream_share=r[5], downstream_share=r[6], yield_rate=r[4],
        method=r[7], evidence="成品罐出料顺序+装车秤+车次记录", confidence="A",
        source="Excel·9-10成品罐和装车批次",
    )

# 补齐缺失的 1→2：入厂原料批次→线边库批次。
# 每第5个线边批次由相邻两个入厂批次汇合，中间入厂批次同时分流至上一组与本组，保证数量守恒。
for i in range(1, 14):
    day = 10 + min((i - 1) // 5, 2)
    batch_id = f"2608{day:02d}-{i:03d}-DZJYM"
    qty = 470 if i == 1 else 30 if i == 13 else 500
    start = datetime(2026, 8, day, (i * 2) % 24, 0)
    add_batch(
        batch_id, order=1, stage="入厂原料批次", material="玉米淀粉", device="原料卸货口",
        start=start, end=start + timedelta(hours=1), qty=qty, status="已验收",
        rule="根据车辆入厂日期、车次和供应商ID划分",
        source="Demo补充·依据批次划分规则",
        note="工作簿缺失层：根据入厂验收、移库与线边收料记录补齐",
    )

for j in range(1, 61):
    target = f"202608{10 + (j - 1) // 24:02d}-{j:03d}-XB01"
    group = (j - 1) // 5 + 1
    day = 10 + min((group - 1) // 5, 2)
    source_a = f"2608{day:02d}-{group:03d}-DZJYM"
    if j % 5:
        add_relation(
            f"REL-F55-12-{j:03d}A", source_a, target, layer="入厂批次→线边库",
            qty=100, upstream_share=100 / batches[source_a]["qty"], downstream_share=1,
            method="WMS移库作业实绩", evidence="入厂验收单+WMS移库单+线边收料确认", confidence="A",
            source="Demo补充·依据批次划分规则",
        )
    else:
        next_group = group + 1
        next_day = 10 + min((next_group - 1) // 5, 2)
        source_b = f"2608{next_day:02d}-{next_group:03d}-DZJYM"
        add_relation(
            f"REL-F55-12-{j:03d}A", source_a, target, layer="入厂批次→线边库",
            qty=70, upstream_share=70 / batches[source_a]["qty"], downstream_share=.7,
            method="WMS移库作业实绩", evidence="入厂验收单+WMS移库单+线边收料确认", confidence="A",
            source="Demo补充·依据批次划分规则",
        )
        add_relation(
            f"REL-F55-12-{j:03d}B", source_b, target, layer="入厂批次→线边库",
            qty=30, upstream_share=30 / batches[source_b]["qty"], downstream_share=.3,
            method="WMS移库作业实绩", evidence="入厂验单+WMS移库单+线边收料确认", confidence="A",
            source="Demo补充·依据批次划分规则",
        )

# 基于麦芽糖 M1/M2/M3 批次规则补充可查询的多对多样例。
malt_source = "Demo补充·麦芽糖批次划分规则"
add_batch(
    "260820-001-JLYM", order=1, stage="入厂原料批次", material="玉米淀粉", device="原料卸货口",
    start="2026-08-20 06:00", end="2026-08-20 08:00", qty=600, line="麦芽糖", route="麦芽糖共用前段",
    status="已验收", rule="根据入厂日期、车次和供应商ID划分", source=malt_source,
)
for i, qty in ((1, 300), (2, 300)):
    xb = f"260820-10{i}-XB0{i}"
    add_batch(
        xb, order=2, stage="线边投料原料批次", material="玉米淀粉", device=f"XB0{i}线边工位",
        start=f"2026-08-20 {7 + i:02d}:00", end=f"2026-08-20 {8 + i:02d}:00", qty=qty, line="麦芽糖", route="麦芽糖共用前段",
        rule="指定移库叉车线边工位，关联平面库或直接卸货的入厂批次", source=malt_source,
    )
    add_relation(
        f"REL-M-12-{i:02d}", "260820-001-JLYM", xb, layer="入厂批次→线边库", qty=qty,
        upstream_share=.5, downstream_share=1, method="WMS移库作业实绩", evidence="入厂验收单+WMS移库单", confidence="A", source=malt_source,
    )

for t, idx in (("0800", 1), ("1200", 2)):
    feed = f"M-260820-S-{t}"
    add_batch(
        feed, order=3, stage="投料批次", material="淀粉乳", device="投料口—淀粉乳接收罐",
        start=f"2026-08-20 {8 + (idx - 1) * 4:02d}:00", end=f"2026-08-20 {12 + (idx - 1) * 4:02d}:00", qty=249,
        line="麦芽糖", route="麦芽糖共用前段", rule="Ts=投料口至淀粉乳接收罐出口有效容积÷质量流量", source=malt_source,
    )
    for x in (1, 2):
        add_relation(
            f"REL-M-23-{idx}{x}", f"260820-10{x}-XB0{x}", feed, layer="线边库→投料", qty=124.5,
            upstream_share=.415, downstream_share=.5, yield_rate=.83, method="投料秤+过程时间窗", evidence="线边投料记录+流量计", confidence="A", source=malt_source,
        )

malt_lines = {
    "M1": ("麦芽糖低DE", "10个糖化罐·128m³", "6个QC罐·53.6m³", "12个成品罐·103m³"),
    "M2": ("麦芽糖中DE", "6个糖化罐·88m³", "3个QC罐·36.4m³", "10个成品罐·76m³"),
    "M3": ("麦芽糖高DE", "6个糖化罐·88m³", "3个QC罐·36.4m³", "10个成品罐·76m³"),
}
for line_idx, (code, info) in enumerate(malt_lines.items(), 1):
    sku, sugar_note, qc_note, fg_note = info
    for t, idx in (("0800", 1), ("1200", 2)):
        wid = f"{code}-260820-W-{t}"
        add_batch(
            wid, order=4, stage="液化批次", material="液化液", device=f"{code}液化线",
            start=f"2026-08-20 {8 + (idx - 1) * 4:02d}:00", end=f"2026-08-20 {12 + (idx - 1) * 4:02d}:00", qty=83,
            line="麦芽糖", route=f"麦芽糖{code}产线", rule="Tw=淀粉乳接收罐出口至糖化罐入口有效容积÷质量流量", source=malt_source,
        )
        add_relation(
            f"REL-{code}-34-{idx}", f"M-260820-S-{t}", wid, layer="投料→液化", qty=83,
            upstream_share=1 / 3, downstream_share=1, method="三条产线支路实际流量", evidence="分流阀位+支路流量计", confidence="A", source=malt_source,
        )
    sugars = []
    for sidx, tank in enumerate(("TH03", "TH04"), 1):
        sid = f"{code}-260820-{tank}-001"
        sugars.append(sid)
        add_batch(
            sid, order=5, stage="糖化批次", material="糖化液", device=tank,
            start=f"2026-08-20 {12 + (sidx - 1) * 4:02d}:00", end=f"2026-08-21 {4 + (sidx - 1) * 4:02d}:00", qty=83,
            line="麦芽糖", route=f"麦芽糖{code}产线", rule="以糖化罐出口阀为物理切分点", source=malt_source, note=sugar_note,
        )
        for idx, t in enumerate(("0800", "1200"), 1):
            add_relation(
                f"REL-{code}-45-{sidx}{idx}", f"{code}-260820-W-{t}", sid, layer="液化→糖化", qty=41.5,
                upstream_share=.5, downstream_share=.5, method="糖化罐进料窗重叠", evidence="液化支路流量+糖化罐进料记录", confidence="B", source=malt_source,
            )
    tid = f"{code}-260821-T-0800"
    add_batch(
        tid, order=6, stage="脱色到蒸发批次", material=sku, device=f"{code}脱色—离交—蒸发",
        start="2026-08-21 08:00", end="2026-08-21 16:00", qty=160,
        line="麦芽糖", route=f"麦芽糖{code}产线", rule="Tt=糖化罐出口至QC罐进口有效容积÷质量流量", source=malt_source,
    )
    for sidx, sid in enumerate(sugars, 1):
        add_relation(
            f"REL-{code}-56-{sidx}", sid, tid, layer="糖化→脱色蒸发", qty=80,
            upstream_share=.964, downstream_share=.5, yield_rate=.964, method="累计流量", evidence="糖化罐出料流量+脱色段入口流量", confidence="A", source=malt_source,
        )
    qids = []
    for qidx in (1, 2):
        qid = f"{code}-260821-Q0{qidx}-01"
        qids.append(qid)
        add_batch(
            qid, order=7, stage="QC罐批次", material=sku, device=f"Q0{qidx}",
            start=f"2026-08-21 {16 + (qidx - 1) * 2:02d}:00", end=f"2026-08-21 {18 + (qidx - 1) * 2:02d}:00", qty=79.6,
            line="麦芽糖", route=f"麦芽糖{code}产线", status="已放行", rule="QC罐独立进料，以进出料阀为物理边界", source=malt_source, note=qc_note,
        )
        add_relation(
            f"REL-{code}-67-{qidx}", tid, qid, layer="脱色蒸发→QC罐", qty=79.6,
            upstream_share=.4975, downstream_share=1, yield_rate=.995, method="QC罐进料累计流量", evidence="DCS阀门+进罐流量+QMS放行单", confidence="A", source=malt_source,
        )
    fid = f"{code}-260822-F001"
    add_batch(
        fid, order=8, stage="成品罐批次", material=sku, device="F001",
        start="2026-08-22 00:00", end="2026-08-22 04:00", qty=158.88,
        line="麦芽糖", route=f"麦芽糖{code}产线", status="已入库", rule="按产品分罐并按进罐时间切分", source=malt_source, note=fg_note,
    )
    for qidx, qid in enumerate(qids, 1):
        add_relation(
            f"REL-{code}-78-{qidx}", qid, fid, layer="QC罐→成品罐", qty=79.44,
            upstream_share=1, downstream_share=.5, yield_rate=.998, method="QC整罐转入", evidence="QC放行+转罐作业记录", confidence="A", source=malt_source,
        )
    for cidx in (1, 2):
        cid = f"{code}-260822-F001-{cidx:02d}"
        add_batch(
            cid, order=9, stage="车次批次", material=sku, device="F001装车位",
            start=f"2026-08-22 {8 + cidx:02d}:00", end=f"2026-08-22 {8 + cidx:02d}:30", qty=25,
            line="麦芽糖", route=f"麦芽糖{code}产线", status="已装车", rule="每车独立，绑定成品罐出料阀和车号", source=malt_source,
        )
        add_relation(
            f"REL-{code}-89-{cidx}", fid, cid, layer="成品罐→装车", qty=25,
            upstream_share=25 / 158.88, downstream_share=1, yield_rate=1, method="装车秤+出料顺序", evidence="成品罐台账+装车秤+车次单", confidence="A", source=malt_source,
        )


# 完整性校验
missing = [r for r in relations if r["from"] not in batches or r["to"] not in batches]
if missing:
    raise RuntimeError(f"存在 {len(missing)} 条关系缺失端点: {missing[:3]}")

layer_counts = {}
for rel in relations:
    layer_counts[rel["layer"]] = layer_counts.get(rel["layer"], 0) + 1

data = {
    "meta": {
        "factory": "中粮太仓工厂",
        "period": "2026-08",
        "snapshot": "2026-09-01 00:35",
        "batchVersion": "BATCH-202608-V2",
        "lineageVersion": "LINEAGE-202608-V2",
        "fruitRouteVersion": "RT-FS-202608-V4",
        "maltRouteVersion": "RT-MALT-202608-V2",
        "status": "已生成·待月度确认",
        "sourceNote": "离线演示数据：果葡糖浆主链采用修正版Excel样例；入厂→线边库与麦芽糖M1/M2/M3为依规则补充。",
    },
    "stats": {
        "batches": len(batches),
        "relations": len(relations),
        "lines": len({b["route"] for b in batches.values()}),
        "layers": layer_counts,
        "excelRelations": sum(1 for r in relations if r["source"].startswith("Excel")),
        "supplementRelations": sum(1 for r in relations if r["source"].startswith("Demo")),
    },
    "batches": sorted(batches.values(), key=lambda x: (x["line"], x["order"], x["start"], x["id"])),
    "relations": relations,
    "rules": {
        "fruit": [
            {"order": 1, "stage": "入厂原料", "format": "YYMMDD-序号-供应商ID", "boundary": "入厂日期+车次+供应商ID"},
            {"order": 2, "stage": "线边库", "format": "YYMMDD-序号-XB工位", "boundary": "指定移库叉车线边工位"},
            {"order": 3, "stage": "投料", "format": "YYMMDD-W-HHMM", "boundary": "Tw=有效容积÷淀粉乳流量"},
            {"order": 4, "stage": "糖化", "format": "YYMMDD-TH罐号-序号", "boundary": "12个糖化罐独立出料"},
            {"order": 5, "stage": "F42中间品", "format": "YYMMDD-F42-HHMM", "boundary": "Tf42=有效容积÷FT120101流量"},
            {"order": 6, "stage": "色谱分离", "format": "YYMMDD-F90-HHMM", "boundary": "Tf90=有效容积÷FT120103流量"},
            {"order": 7, "stage": "F55混合", "format": "YYMMDD-H-HHMM", "boundary": "两路汇合+配方切换+Th时间窗"},
            {"order": 8, "stage": "QC罐", "format": "YYMMDD-Q罐号-进料序号", "boundary": "QC罐进出料阀"},
            {"order": 9, "stage": "成品罐", "format": "YYMMDD-F罐号", "boundary": "按产品分罐+进罐时间"},
            {"order": 10, "stage": "车次", "format": "YYMMDD-F罐号-装车序号", "boundary": "每车独立+车号绑定"},
        ],
        "malt": [
            {"order": 1, "stage": "入厂原料", "format": "YYMMDD-序号-供应商ID", "boundary": "入厂日期+车次+供应商ID"},
            {"order": 2, "stage": "线边投料原料", "format": "YYMMDD-序号-XB工位", "boundary": "指定移库叉车线边工位"},
            {"order": 3, "stage": "投料", "format": "M-YYMMDD-S-HHMM", "boundary": "Ts=有效容积÷淀粉乳流量"},
            {"order": 4, "stage": "液化", "format": "M1/M2/M3-YYMMDD-W-HHMM", "boundary": "Tw=有效容积÷淀粉乳流量"},
            {"order": 5, "stage": "糖化", "format": "M1/M2/M3-YYMMDD-TH罐号-序号", "boundary": "糖化罐独立出料"},
            {"order": 6, "stage": "脱色到蒸发", "format": "M1/M2/M3-YYMMDD-T-HHMM", "boundary": "Tt=有效容积÷质量流量"},
            {"order": 7, "stage": "QC罐", "format": "M1/M2/M3-YYMMDD-Q罐号-进料序号", "boundary": "QC罐进出料阀"},
            {"order": 8, "stage": "成品罐", "format": "M1/M2/M3-YYMMDD-F罐号", "boundary": "按产品分罐+进罐时间"},
            {"order": 9, "stage": "车次", "format": "M1/M2/M3-YYMMDD-F罐号-装车序号", "boundary": "每车独立+车号绑定"},
        ],
    },
}

json_text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
OUTPUT_JS.parent.mkdir(parents=True, exist_ok=True)
OUTPUT_JSON.write_text(json_text + "\n", encoding="utf-8")
OUTPUT_JS.write_text("window.COSTING_LINEAGE_DATA=" + json_text + ";\n", encoding="utf-8")
print(json.dumps({"batches": len(batches), "relations": len(relations), "layers": layer_counts}, ensure_ascii=False, indent=2))
