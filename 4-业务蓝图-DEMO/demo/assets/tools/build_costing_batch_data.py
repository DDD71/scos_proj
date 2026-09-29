#!/usr/bin/env python3
"""构建批次直接成本、费用池分摊和批次成本离线演示数据。"""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
DEMO_DIR = HERE.parents[1]
BLUEPRINT_DIR = DEMO_DIR.parents[2]
SOURCE_BOOK = BLUEPRINT_DIR / "成本分摊模块" / "费用池清单_简化版.xlsx"
LINEAGE_JSON = DEMO_DIR / "assets" / "data" / "costing" / "costing-lineage-data.json"
OUTPUT_JSON = DEMO_DIR / "assets" / "data" / "costing" / "costing-batch-data-v3.json"
OUTPUT_JS = DEMO_DIR / "assets" / "scripts" / "costing" / "costing-batch-data-v3.js"


OPERATING_CATEGORIES = {"管理费用", "销售费用", "财务费用", "税费及其他"}
PRODUCTION_STAGES = {
    "投料批次",
    "液化批次",
    "糖化批次",
    "F42中间品批次",
    "色谱分离批次",
    "F55混合批次",
    "脱色到蒸发批次",
    "QC罐批次",
    "成品罐批次",
}
PROCESS_STAGES = PRODUCTION_STAGES - {"QC罐批次", "成品罐批次"}
INBOUND_STAGES = {"入厂原料批次", "线边库批次", "线边投料原料批次"}
OUTBOUND_STAGES = {"车次批次"}
DIRECT_AUXILIARY_SUBCATEGORIES = {"酸碱及过滤辅料", "酶制剂及配方辅料"}
PERIODIC_AUXILIARY_CODES = {"CE016", "CE024", "CE025", "CE026"}


def pool_rule(name, scope, primary, backup, stages, unit, allocate=True, view="生产制造成本"):
    return {
        "name": name,
        "scope": scope,
        "primaryDriver": primary,
        "backupDriver": backup,
        "stages": sorted(stages),
        "driverUnit": unit,
        "allocateToBatch": allocate,
        "costView": view,
    }


POOL_RULES = {
    "POOL_RAW_PRICE_DIFF": pool_rule("原料价差池", "入厂原料批次", "实际入厂质量", "采购收货质量", {"入厂原料批次"}, "吨"),
    "POOL_PROCUREMENT_INS": pool_rule("采购附加费用池", "采购批与入厂原料批次", "采购批金额", "实际入厂质量", {"入厂原料批次"}, "万元"),
    "POOL_EXTERNAL_WH": pool_rule("外库成本池", "发生外库占用的原料批次", "批次外库占用吨天", "实际入库质量", {"入厂原料批次", "线边库批次"}, "吨天"),
    "POOL_INBOUND_LOG": pool_rule("入厂物流作业池", "入厂、移库与装车作业批次", "车次作业量", "实际处理质量", INBOUND_STAGES | OUTBOUND_STAGES, "次"),
    "POOL_WAREHOUSE_LABOR": pool_rule("仓储投料人工池", "仓储和投料受益批次", "批次有效工时", "实际处理量", {"线边库批次", "线边投料原料批次", "投料批次"}, "小时"),
    "POOL_PROCESS_AUX": pool_rule("工艺辅料池", "使用酸碱及过滤辅料的工艺批次", "批次实际处理量", "绝干质量", PROCESS_STAGES, "吨"),
    "POOL_ENZYME": pool_rule("酶制剂池", "液化、糖化和异构相关批次", "实际投加或处理量", "批次有效工时", {"投料批次", "液化批次", "糖化批次", "F42中间品批次", "色谱分离批次", "F55混合批次"}, "吨"),
    "POOL_LONG_CYCLE": pool_rule("长周期材料池", "树脂及循环材料实际服务批次", "寿命单位处理量", "设备运行小时", {"F42中间品批次", "色谱分离批次", "F55混合批次", "脱色到蒸发批次"}, "吨"),
    "POOL_PACKAGING": pool_rule("包装物费用池", "成品与发运批次", "扫码领用数量", "成品或发运质量", {"成品罐批次", "车次批次"}, "吨"),
    "POOL_WATER": pool_rule("公共水费用池", "公共水边界内生产批次", "分表实测耗水", "处理量或绝干量", PRODUCTION_STAGES, "吨"),
    "POOL_ELECTRICITY": pool_rule("公共电力费用池", "公共电表边界内生产批次", "分表实测电量", "设备运行小时", PRODUCTION_STAGES, "千瓦时"),
    "POOL_STEAM": pool_rule("公共蒸汽费用池", "蒸汽边界内生产批次", "分表实测蒸汽量", "处理量或绝干量", PROCESS_STAGES, "吨"),
    "POOL_NATURAL_GAS": pool_rule("天然气费用池", "燃气设备服务批次", "分表实测气量", "设备运行小时", PROCESS_STAGES, "立方米"),
    "POOL_BIOGAS": pool_rule("沼气费用池", "使用沼气的生产批次", "分表实测气量", "设备运行小时", PROCESS_STAGES, "立方米"),
    "POOL_OTHER_UTILITY": pool_rule("其他动力费用池", "公用工程服务批次", "批次专属实测消耗", "处理量或绝干量", PRODUCTION_STAGES, "吨"),
    "POOL_WW_VARIABLE": pool_rule("污水处理变动成本池", "产生排放的工艺批次", "污水量×污染系数", "投料量或产量", PROCESS_STAGES, "污染当量"),
    "POOL_WW_FIXED": pool_rule("污水处理固定成本池", "污水设施受益产线", "设施服务量或运行小时", "绝干量", PROCESS_STAGES, "小时"),
    "POOL_WASTE_DISPOSAL": pool_rule("固废危废处置池", "产生固废或危废的工艺批次", "处置单来源质量", "实际处理量", PROCESS_STAGES, "吨"),
    "POOL_QUALITY": pool_rule("质量检验费用池", "QC及被检生产批次", "检验项目数", "样品数或被检批次质量", {"QC罐批次"}, "项"),
    "POOL_REWORK_LOSS": pool_rule("返工报废损失池", "经确认的异常责任对象", "责任批实际损失量", "返工作业工时", set(), "吨", False),
    "POOL_BYPRODUCT_OFFSET": pool_rule("副产及回收收益抵减池", "产生副产物的工艺批次", "副产物实际产出量", "主产品实际产量", PROCESS_STAGES, "吨"),
    "POOL_PROD_LABOR": pool_rule("生产人员人工池", "班组服务的产线和节点", "批次有效工时", "处理量或标准工时", PRODUCTION_STAGES, "小时"),
    "POOL_OUTSOURCE_LABOR": pool_rule("生产外包劳务池", "外包班组服务批次", "批次有效工时", "处理量或标准工时", PRODUCTION_STAGES, "小时"),
    "POOL_DEPRECIATION": pool_rule("生产设备折旧池", "设备实际服务批次", "设备占用或运行小时", "设备处理量", PRODUCTION_STAGES, "小时"),
    "POOL_AMORTIZATION": pool_rule("生产资产摊销池", "资产实际服务批次", "资产受益小时", "设备处理量", PRODUCTION_STAGES, "小时"),
    "POOL_MAINTENANCE": pool_rule("维修保养费用池", "受维修设备服务批次", "维修工单影响小时", "设备运行小时", PRODUCTION_STAGES, "小时"),
    "POOL_MFG_OVERHEAD": pool_rule("公共制造费用池", "对应车间和装置范围", "经确认的复合动因", "绝干量", PRODUCTION_STAGES, "标准动因"),
    "POOL_PROD_INS": pool_rule("生产保险费用池", "受保资产服务批次", "资产原值×受益时间", "设备小时或绝干量", PRODUCTION_STAGES, "受益点"),
    "POOL_IDLE": pool_rule("停工及闲置成本池", "停工或闲置成本对象", "停工对象实际维持成本", "无", set(), "元", False),
    "POOL_TANKER_LOADING": pool_rule("槽车及装车费用池", "销售发运批次", "车次作业量", "发运质量", OUTBOUND_STAGES, "车次"),
    "POOL_OUTBOUND_TRANS": pool_rule("客户运输费用池", "销售发运批次", "吨公里", "发运质量", OUTBOUND_STAGES, "吨公里"),
    "POOL_OUTBOUND_INS": pool_rule("货运保险费用池", "保单覆盖的发运批次", "保单覆盖质量", "发运质量", OUTBOUND_STAGES, "吨"),
}

for code, name in {
    "POOL_MANAGEMENT": "管理费用池",
    "POOL_SELLING": "销售费用池",
    "POOL_FINANCE": "财务费用池",
    "POOL_TAX_OTHER": "税费及其他费用池",
}.items():
    POOL_RULES[code] = pool_rule(name, "完全经营成本视图", "收入或批准的复合动因", "销量", set(), "元", False, "完全经营成本")


def read_cost_items():
    ws = load_workbook(SOURCE_BOOK, read_only=True, data_only=True).active
    items = []
    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        code = str(row[0]).strip()
        if not code.startswith("CE"):
            continue
        category = str(row[1] or "").strip()
        treatment = str(row[4] or "").strip()
        item = {
            "code": code,
            "category": category,
            "subcategory": str(row[2] or "").strip(),
            "name": str(row[3] or "").strip(),
            "treatment": treatment,
            "poolCode": str(row[5] or "").strip(),
            "directCondition": str(row[6] or "").strip(),
            "poolCondition": str(row[7] or "").strip(),
            "sourceSystem": str(row[8] or "").strip(),
            "costView": "完全经营成本" if category in OPERATING_CATEGORIES else "生产制造成本",
            "directEligibility": "默认直计" if treatment == "通常直接计批" else ("满足条件可直计" if category not in OPERATING_CATEGORIES else "不计生产批次"),
        }
        if code in PERIODIC_AUXILIARY_CODES:
            item.update({
                "treatment": "周期费用池",
                "poolCode": "POOL_LONG_CYCLE",
                "directCondition": "周期性更换材料跨多个批次服务，不直接计入单一批次",
                "poolCondition": "按材料寿命周期覆盖的实际服务批次分摊",
                "directEligibility": "不可直接计批",
            })
        elif item["subcategory"] in DIRECT_AUXILIARY_SUBCATEGORIES:
            item.update({
                "treatment": "通常直接计批",
                "poolCode": "",
                "directCondition": "按批次领用、配方投加或工序投加记录直接归集",
                "poolCondition": "不进入费用池",
                "directEligibility": "默认直计",
            })
        items.append(item)
    return items


def product_category(batch):
    return "果糖类" if batch["line"] == "果葡糖浆" else "麦芽糖类"


def duration_hours(batch):
    try:
        start = datetime.strptime(batch["start"], "%Y-%m-%d %H:%M")
        end = datetime.strptime(batch["end"], "%Y-%m-%d %H:%M")
        return max((end - start).total_seconds() / 3600, 0.5)
    except (KeyError, TypeError, ValueError):
        return 1.0


def direct_specs(batch, ordinal):
    qty = float(batch.get("qty") or 0)
    stage = batch["stage"]
    specs = {
        "入厂原料批次": [("CE001", qty, "吨", 2680, "采购收货单"), ("CE002", qty, "吨", 48, "运输结算单")],
        "线边库批次": [("CE010", qty, "吨", 8, "移库作业单")],
        "线边投料原料批次": [("CE010", qty, "吨", 8, "移库作业单")],
        "投料批次": [
            ("CE013", qty * 0.006, "吨", 1100, "MES投加记录"),
            ("CE014", qty * 0.004, "吨", 3200, "MES投加记录"),
            ("CE018", qty * 0.0025, "吨", 48000, "配方投加记录"),
            ("CE021", qty * 0.0006, "吨", 4600, "配方投加记录"),
            ("CE022", qty * 0.0004, "吨", 4200, "配方投加记录"),
        ],
        "液化批次": [
            ("CE018", qty * 0.0022, "吨", 48000, "配方投加记录"),
            ("CE022", qty * 0.00035, "吨", 4200, "配方投加记录"),
            ("CE023", qty * 0.0005, "吨", 6800, "工序投加记录"),
            ("CE031", qty * 0.24, "吨", 235, "蒸汽分表"),
        ],
        "糖化批次": [
            ("CE019", qty * 0.0018, "吨", 42000, "配方投加记录"),
            ("CE021", qty * 0.0005, "吨", 4600, "配方投加记录"),
            ("CE023", qty * 0.0004, "吨", 6800, "工序投加记录"),
            ("CE030", qty * 35, "千瓦时", 0.72, "电力分表"),
        ],
        "F42中间品批次": [
            ("CE015", qty * 0.003, "吨", 3500, "过滤辅料领用记录"),
            ("CE017", qty * 0.0015, "吨", 12000, "过滤耗材领用记录"),
            ("CE023", qty * 0.0004, "吨", 6800, "工序投加记录"),
            ("CE030", qty * 42, "千瓦时", 0.72, "电力分表"),
        ],
        "色谱分离批次": [
            ("CE020", qty * 0.0012, "吨", 52000, "配方投加记录"),
            ("CE021", qty * 0.0005, "吨", 4600, "配方投加记录"),
            ("CE030", qty * 75, "千瓦时", 0.72, "电力分表"),
        ],
        "F55混合批次": [
            ("CE020", qty * 0.0008, "吨", 52000, "配方投加记录"),
            ("CE022", qty * 0.0003, "吨", 4200, "配方投加记录"),
            ("CE031", qty * 0.16, "吨", 235, "蒸汽分表"),
        ],
        "脱色到蒸发批次": [
            ("CE015", qty * 0.0028, "吨", 3500, "过滤辅料领用记录"),
            ("CE017", qty * 0.0012, "吨", 12000, "过滤耗材领用记录"),
            ("CE023", qty * 0.00035, "吨", 6800, "工序投加记录"),
            ("CE030", qty * 55, "千瓦时", 0.72, "电力分表"),
            ("CE031", qty * 0.12, "吨", 235, "蒸汽分表"),
        ],
        "QC罐批次": [("CE038", 3 + ordinal % 3, "项", 380, "LIMS检验单")],
        "成品罐批次": [("CE029", qty * 0.08, "吨", 3.5, "水表区间"), ("CE030", qty * 8, "千瓦时", 0.72, "电力分表")],
        "车次批次": [("CE053", 1, "车次", 650, "装车作业单"), ("CE055", qty, "吨", 65, "运输结算单"), ("CE056", qty, "吨", 4, "货运保单")],
    }.get(stage, [])
    if stage == "QC罐批次" and ordinal % 4 == 0:
        specs.append(("CE039", 1, "项", 1800, "外部检测结算单"))
    return specs


def pool_amount(item):
    number = int(item["code"][2:])
    base = {
        "原料采购及入厂": 36000,
        "生产变动成本": 52000,
        "生产固定成本": 180000,
        "出库及销售物流": 46000,
        "管理费用": 150000,
        "销售费用": 120000,
        "财务费用": 90000,
        "税费及其他": 70000,
    }[item["category"]]
    amount = round(base * (0.72 + (number % 7) * 0.09), 2)
    if item["treatment"] == "收益抵减池":
        amount = -amount
    return amount


def driver_value(batch, rule, ordinal):
    qty = float(batch.get("qty") or 0)
    primary = rule["primaryDriver"]
    if "吨公里" in primary:
        return qty * (120 + (ordinal % 5) * 35)
    if "工时" in primary or "小时" in primary or "受益时间" in primary:
        return duration_hours(batch)
    if "车次" in primary or "作业量" in primary:
        return 1.0
    if "检验" in primary or "样品" in primary:
        return float(3 + ordinal % 3)
    if "污染" in primary:
        return qty * (0.72 + (ordinal % 4) * 0.08)
    if "金额" in primary:
        return max(qty * 2680 / 10000, 0.01)
    if "复合动因" in primary or "受益点" in rule["driverUnit"]:
        return qty * (1 + (ordinal % 3) * 0.05)
    return max(qty, 0.01)


cost_items = read_cost_items()
item_by_code = {item["code"]: item for item in cost_items}
lineage = json.loads(LINEAGE_JSON.read_text(encoding="utf-8"))
batches = lineage["batches"]
relations = lineage["relations"]
batch_by_id = {batch["id"]: batch for batch in batches}

direct_details = []
direct_by_batch = defaultdict(list)
for ordinal, batch in enumerate(batches):
    for item_code, qty, unit, unit_price, evidence in direct_specs(batch, ordinal):
        item = item_by_code[item_code]
        amount = round(qty * unit_price, 2)
        row = {
            "id": f"DIR-{batch['id']}-{item_code}",
            "batch": batch["id"],
            "productCategory": product_category(batch),
            "line": batch["line"],
            "stage": batch["stage"],
            "material": batch["material"],
            "itemCode": item_code,
            "costCategory": item["category"],
            "costSubcategory": item["subcategory"],
            "costName": item["name"],
            "quantity": round(qty, 4),
            "unit": unit,
            "unitPrice": round(unit_price, 4),
            "amount": amount,
            "source": item["sourceSystem"],
            "evidence": evidence,
            "status": "已归集",
        }
        direct_details.append(row)
        direct_by_batch[batch["id"]].append(row)

direct_summaries = []
for batch in batches:
    total = round(sum(row["amount"] for row in direct_by_batch[batch["id"]]), 2)
    direct_summaries.append({
        "batch": batch["id"],
        "productCategory": product_category(batch),
        "line": batch["line"],
        "stage": batch["stage"],
        "material": batch["material"],
        "quantity": batch["qty"],
        "unit": batch["unit"],
        "detailCount": len(direct_by_batch[batch["id"]]),
        "directTotal": total,
        "status": "已归集" if total else "本节点无新增直计成本",
    })

pool_cost_lines = []
pool_groups = defaultdict(list)
for item in cost_items:
    if not item["poolCode"]:
        continue
    rule = POOL_RULES[item["poolCode"]]
    amount = pool_amount(item)
    row = {
        **item,
        "poolName": rule["name"],
        "benefitScope": rule["scope"],
        "primaryDriver": rule["primaryDriver"],
        "backupDriver": rule["backupDriver"],
        "periodAmount": amount,
        "allocationStatus": "待分摊",
    }
    pool_cost_lines.append(row)
    pool_groups[item["poolCode"]].append(row)

pool_allocations = []
pool_alloc_by_batch = defaultdict(list)
for pool_code, lines in pool_groups.items():
    rule = POOL_RULES[pool_code]
    if not rule["allocateToBatch"]:
        status = "完全成本视图" if rule["costView"] == "完全经营成本" else "独立成本对象"
        for line in lines:
            line["allocationStatus"] = status
        continue
    eligible = [batch for batch in batches if batch["stage"] in set(rule["stages"])]
    drivers = [(batch, driver_value(batch, rule, ordinal)) for ordinal, batch in enumerate(eligible)]
    denominator = sum(value for _, value in drivers)
    if not eligible or denominator <= 0:
        for line in lines:
            line["allocationStatus"] = "未分摊"
        continue
    for line in lines:
        allocated = 0.0
        rate = line["periodAmount"] / denominator
        for index, (batch, driver_qty) in enumerate(drivers):
            amount = round(rate * driver_qty, 2)
            if index == len(drivers) - 1:
                amount = round(line["periodAmount"] - allocated, 2)
            allocated += amount
            row = {
                "id": f"ALLOC-{line['code']}-{batch['id']}",
                "poolCode": pool_code,
                "poolName": rule["name"],
                "itemCode": line["code"],
                "costCategory": line["category"],
                "costName": line["name"],
                "batch": batch["id"],
                "productCategory": product_category(batch),
                "line": batch["line"],
                "stage": batch["stage"],
                "material": batch["material"],
                "driverName": rule["primaryDriver"],
                "driverQuantity": round(driver_qty, 4),
                "driverUnit": rule["driverUnit"],
                "allocationRate": round(rate, 6),
                "amount": amount,
                "confidence": "A",
                "status": "已分摊",
            }
            pool_allocations.append(row)
            pool_alloc_by_batch[batch["id"]].append(row)
        line["allocationStatus"] = "已分摊" if line["periodAmount"] >= 0 else "已抵减"

pool_definitions = []
for pool_code, lines in sorted(pool_groups.items()):
    rule = POOL_RULES[pool_code]
    total = round(sum(line["periodAmount"] for line in lines), 2)
    statuses = {line["allocationStatus"] for line in lines}
    status = "已分摊" if statuses <= {"已分摊", "已抵减"} else sorted(statuses)[0]
    pool_definitions.append({
        "poolCode": pool_code,
        "poolName": rule["name"],
        "costView": rule["costView"],
        "benefitScope": rule["scope"],
        "primaryDriver": rule["primaryDriver"],
        "backupDriver": rule["backupDriver"],
        "itemCount": len(lines),
        "periodAmount": total,
        "status": status,
    })

pool_summaries = []
for batch in batches:
    rows = pool_alloc_by_batch[batch["id"]]
    total = round(sum(row["amount"] for row in rows), 2)
    pool_summaries.append({
        "batch": batch["id"],
        "productCategory": product_category(batch),
        "line": batch["line"],
        "stage": batch["stage"],
        "material": batch["material"],
        "quantity": batch["qty"],
        "unit": batch["unit"],
        "poolCount": len({row["poolCode"] for row in rows}),
        "detailCount": len(rows),
        "poolTotal": total,
        "status": "已分摊" if rows else "无批次分摊",
    })

incoming = defaultdict(list)
for relation in relations:
    incoming[relation["to"]].append(relation)

batch_costs = []
batch_cost_by_id = {}
batch_cost_details = []
for batch in sorted(batches, key=lambda row: (row["order"], row["start"], row["id"])):
    inbound_total = 0.0
    for relation in incoming[batch["id"]]:
        upstream = batch_cost_by_id[relation["from"]]
        upstream_batch = batch_by_id[relation["from"]]
        transfer_qty = float(relation.get("qty") or 0)
        ratio = transfer_qty / float(upstream_batch.get("qty") or 1)
        amount = round(upstream["totalCost"] * ratio, 2)
        inbound_total += amount
        batch_cost_details.append({
            "batch": batch["id"],
            "component": "上游成本转入",
            "costCategory": "上游成本转入",
            "costName": f"{upstream_batch['stage']} · {relation['from']}",
            "productCategory": product_category(batch),
            "stage": batch["stage"],
            "quantity": transfer_qty,
            "unit": relation.get("unit") or "吨",
            "unitPrice": round(amount / transfer_qty, 4) if transfer_qty else 0,
            "amount": amount,
            "source": f"{relation['id']} · {relation['method']}",
        })
    direct_total = round(sum(row["amount"] for row in direct_by_batch[batch["id"]]), 2)
    pool_total = round(sum(row["amount"] for row in pool_alloc_by_batch[batch["id"]]), 2)
    adjustment = 0.0
    total = round(inbound_total + direct_total + pool_total + adjustment, 2)
    unit_cost = round(total / float(batch.get("qty") or 1), 2)
    result = {
        "batch": batch["id"],
        "productCategory": product_category(batch),
        "line": batch["line"],
        "stage": batch["stage"],
        "material": batch["material"],
        "quantity": batch["qty"],
        "unit": batch["unit"],
        "inboundCost": round(inbound_total, 2),
        "directCost": direct_total,
        "poolCost": pool_total,
        "adjustment": adjustment,
        "totalCost": total,
        "unitCost": unit_cost,
        "status": "财务锁定",
    }
    batch_costs.append(result)
    batch_cost_by_id[batch["id"]] = result
    for row in direct_by_batch[batch["id"]]:
        batch_cost_details.append({
            "batch": batch["id"],
            "component": "批次直接成本",
            "costCategory": row["costCategory"],
            "costName": row["costName"],
            "productCategory": row["productCategory"],
            "stage": row["stage"],
            "quantity": row["quantity"],
            "unit": row["unit"],
            "unitPrice": row["unitPrice"],
            "amount": row["amount"],
            "source": row["evidence"],
        })
    for row in pool_alloc_by_batch[batch["id"]]:
        batch_cost_details.append({
            "batch": batch["id"],
            "component": "费用池分摊成本",
            "costCategory": row["poolName"],
            "costName": row["costName"],
            "productCategory": row["productCategory"],
            "stage": row["stage"],
            "quantity": row["driverQuantity"],
            "unit": row["driverUnit"],
            "unitPrice": row["allocationRate"],
            "amount": row["amount"],
            "source": f"{row['poolCode']} · {row['driverName']}",
        })

batch_costs.sort(key=lambda row: (row["line"], row["stage"], row["batch"]))
direct_summaries.sort(key=lambda row: (row["line"], row["stage"], row["batch"]))
pool_summaries.sort(key=lambda row: (row["line"], row["stage"], row["batch"]))

allocated_pool_total = round(sum(row["amount"] for row in pool_allocations), 2)
manufacturing_pool_total = round(sum(row["periodAmount"] for row in pool_cost_lines if POOL_RULES[row["poolCode"]]["allocateToBatch"]), 2)
direct_total = round(sum(row["amount"] for row in direct_details), 2)

data = {
    "meta": {
        "factory": "中粮太仓工厂",
        "period": lineage["meta"]["period"],
        "snapshot": lineage["meta"]["snapshot"],
        "costVersion": "COST-202608-V4",
        "poolRuleVersion": "POOLRULE-202608-V3",
        "sourceWorkbook": SOURCE_BOOK.name,
        "lineageVersion": lineage["meta"]["lineageVersion"],
    },
    "stats": {
        "costItems": len(cost_items),
        "directCatalogItems": sum(1 for item in cost_items if item["category"] not in OPERATING_CATEGORIES and item["directEligibility"] != "不可直接计批"),
        "operatingItems": sum(1 for item in cost_items if item["category"] in OPERATING_CATEGORIES),
        "poolCostLines": len(pool_cost_lines),
        "poolCount": len(pool_definitions),
        "batches": len(batches),
        "directDetails": len(direct_details),
        "poolAllocations": len(pool_allocations),
        "directTotal": direct_total,
        "manufacturingPoolTotal": manufacturing_pool_total,
        "allocatedPoolTotal": allocated_pool_total,
        "poolReconciliationDifference": round(manufacturing_pool_total - allocated_pool_total, 2),
    },
    "directCatalog": [item for item in cost_items if item["category"] not in OPERATING_CATEGORIES and item["directEligibility"] != "不可直接计批"],
    "directDetails": direct_details,
    "directSummaries": direct_summaries,
    "poolCostLines": pool_cost_lines,
    "poolDefinitions": pool_definitions,
    "poolAllocations": pool_allocations,
    "poolSummaries": pool_summaries,
    "batchCosts": batch_costs,
    "batchCostDetails": batch_cost_details,
}

json_text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
OUTPUT_JS.parent.mkdir(parents=True, exist_ok=True)
OUTPUT_JSON.write_text(json_text + "\n", encoding="utf-8")
OUTPUT_JS.write_text("window.COSTING_BATCH_DATA_V3=" + json_text + ";\n", encoding="utf-8")
print(json.dumps(data["stats"], ensure_ascii=False, indent=2))
